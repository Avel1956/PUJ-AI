"""rag_engine.py — Motor RAG con LangChain + ChromaDB (ONNX, sin torch)."""
import hashlib
import os

from langchain_text_splitters import RecursiveCharacterTextSplitter
from langchain_community.document_loaders import PyPDFLoader
from langchain_chroma import Chroma
from langchain_core.documents import Document
from chromadb.utils import embedding_functions


BASE_DIR = os.path.dirname(os.path.abspath(__file__))
ASIGNATURAS_DIR = os.path.join(BASE_DIR, "asignaturas")
VECTORSTORE_DIR = os.path.join(BASE_DIR, ".vectorstores")

# Formatos soportados. Antes solo se leían .pdf/.txt/.md, así que cualquier
# .docx o .pptx colocado en documentos/ se ignoraba en silencio.
EXTENSIONES_TEXTO = (".txt", ".md")
EXTENSIONES_SOPORTADAS = (".pdf",) + EXTENSIONES_TEXTO + (".docx", ".pptx")

CHUNK_SIZE = 800
CHUNK_OVERLAP = 150


# ============================================================
# Extracción de texto por formato
# ============================================================
def _extraer_docx(ruta: str) -> str:
    """Texto de un .docx, incluido el contenido de las tablas."""
    from docx import Document as DocxDocument

    documento = DocxDocument(ruta)
    partes = [p.text.strip() for p in documento.paragraphs if p.text and p.text.strip()]
    for tabla in documento.tables:
        for fila in tabla.rows:
            celdas = [c.text.strip() for c in fila.cells]
            if any(celdas):
                partes.append(" | ".join(celdas))
    return "\n\n".join(partes)


def _extraer_pptx(ruta: str) -> str:
    """Texto de un .pptx: cuadros de texto, tablas y notas del ponente."""
    from pptx import Presentation

    presentacion = Presentation(ruta)
    partes = []
    for i, diapositiva in enumerate(presentacion.slides, 1):
        partes.append(f"## Diapositiva {i}")
        for forma in diapositiva.shapes:
            if getattr(forma, "has_text_frame", False) and forma.text_frame.text.strip():
                partes.append(forma.text_frame.text.strip())
            if getattr(forma, "has_table", False) and forma.has_table:
                for fila in forma.table.rows:
                    celdas = [c.text.strip() for c in fila.cells]
                    if any(celdas):
                        partes.append(" | ".join(celdas))
        # Las notas del ponente suelen contener la explicación real del tema.
        notas = getattr(diapositiva, "notes_slide", None)
        if notas is not None:
            try:
                texto_notas = notas.notes_text_frame.text.strip()
                if texto_notas:
                    partes.append(f"Notas del ponente: {texto_notas}")
            except Exception:
                pass
    return "\n\n".join(partes)


def _leer_texto(ruta: str) -> str:
    """Lee .txt/.md probando varias codificaciones antes de rendirse.

    Antes se forzaba utf-8: un archivo en cualquier otra codificación fallaba
    y se descartaba sin aviso.
    """
    for codec in ("utf-8", "utf-8-sig", "cp1252", "latin-1"):
        try:
            with open(ruta, "r", encoding=codec) as f:
                return f.read()
        except (UnicodeDecodeError, OSError):
            continue
    return ""


class ChromaEmbedder:
    """Wrapper ligero sobre ChromaDB ONNX — sin torch ni sentence-transformers."""

    def __init__(self):
        self._ef = embedding_functions.DefaultEmbeddingFunction()

    def embed_documents(self, texts: list[str]) -> list[list[float]]:
        return self._ef(texts)

    def embed_query(self, text: str) -> list[float]:
        return self._ef([text])[0]


class MotorRAG:
    """Carga documentos de una asignatura, indexa en ChromaDB, y provee retrieval."""

    def __init__(self, nombre_asignatura: str):
        self.nombre = nombre_asignatura
        self.ruta = os.path.join(ASIGNATURAS_DIR, nombre_asignatura)
        self.ruta_docs = os.path.join(self.ruta, "documentos")
        self.persist_dir = os.path.join(VECTORSTORE_DIR, nombre_asignatura)
        self.ruta_fingerprint = os.path.join(self.persist_dir, "_fingerprint.txt")

        self.embeddings = ChromaEmbedder()

        self.documents: list[Document] = []
        self.vectorstore: Chroma | None = None
        self._fingerprint: str = ""
        self.archivos_ignorados: list[str] = []

    # ----- Carga de documentos -----
    def _listar_archivos(self) -> list[tuple[str, str]]:
        """Recorre documentos/ de forma recursiva. Devuelve [(ruta, relativa)]."""
        encontrados = []
        if not os.path.isdir(self.ruta_docs):
            return encontrados
        for raiz, _dirs, archivos in os.walk(self.ruta_docs):
            for archivo in sorted(archivos):
                if archivo.startswith("."):
                    continue  # .gitkeep y similares
                ruta_completa = os.path.join(raiz, archivo)
                relativa = os.path.relpath(ruta_completa, self.ruta_docs)
                # Se normaliza a "/" para que la etiqueta y el fingerprint sean
                # idénticos en Windows y en Linux.
                encontrados.append((ruta_completa, relativa.replace(os.sep, "/")))
        return sorted(encontrados, key=lambda x: x[1])

    def cargar_documentos(self) -> int:
        """Carga PDF/Word/PowerPoint/txt/md desde documentos/. Retorna cuántos docs."""
        self.documents = []
        self.archivos_ignorados = []

        for ruta_completa, relativa in self._listar_archivos():
            ext = os.path.splitext(relativa)[1].lower()

            if ext not in EXTENSIONES_SOPORTADAS:
                self.archivos_ignorados.append(f"{relativa} (formato no soportado)")
                continue
            try:
                if ext == ".pdf":
                    self.documents.extend(PyPDFLoader(ruta_completa).load())
                elif ext in EXTENSIONES_TEXTO:
                    texto = _leer_texto(ruta_completa)
                    if texto.strip():
                        self.documents.append(
                            Document(page_content=texto, metadata={"source": relativa})
                        )
                    else:
                        self.archivos_ignorados.append(f"{relativa} (vacío o ilegible)")
                elif ext == ".docx":
                    texto = _extraer_docx(ruta_completa)
                    if texto.strip():
                        self.documents.append(
                            Document(page_content=texto, metadata={"source": relativa})
                        )
                    else:
                        self.archivos_ignorados.append(f"{relativa} (sin texto extraíble)")
                elif ext == ".pptx":
                    texto = _extraer_pptx(ruta_completa)
                    if texto.strip():
                        self.documents.append(
                            Document(page_content=texto, metadata={"source": relativa})
                        )
                    else:
                        self.archivos_ignorados.append(f"{relativa} (sin texto extraíble)")
            except ImportError as e:
                self.archivos_ignorados.append(f"{relativa} (falta dependencia: {e.name})")
            except Exception as e:
                self.archivos_ignorados.append(f"{relativa} ({type(e).__name__})")

        return len(self.documents)

    # ----- Fingerprint para cache -----
    def _calcular_fingerprint(self) -> str:
        """Hash del contenido de documentos/ (recursivo) para detectar cambios."""
        if not os.path.isdir(self.ruta_docs):
            return ""
        hasher = hashlib.md5()
        for ruta_completa, relativa in self._listar_archivos():
            hasher.update(relativa.encode())
            try:
                hasher.update(str(os.path.getmtime(ruta_completa)).encode())
                hasher.update(str(os.path.getsize(ruta_completa)).encode())
            except OSError:
                pass
        return hasher.hexdigest()

    def _leer_fingerprint_persistido(self) -> str:
        try:
            with open(self.ruta_fingerprint, "r", encoding="utf-8") as f:
                return f.read().strip()
        except OSError:
            return ""

    def _escribir_fingerprint(self, fp: str) -> None:
        try:
            os.makedirs(self.persist_dir, exist_ok=True)
            with open(self.ruta_fingerprint, "w", encoding="utf-8") as f:
                f.write(fp)
        except OSError:
            pass

    def _nombre_coleccion(self, fp: str) -> str:
        """Nombre de colección derivado del fingerprint.

        Es la clave para no acumular duplicados: cuando el contenido cambia,
        el fingerprint cambia y la colección resultante es OTRA, así que el
        índice viejo deja de consultarse en lugar de mezclarse con el nuevo.
        """
        return f"{self.nombre}-{(fp or 'vacio')[:12]}"

    # ----- Indexación -----
    def indexar(self, force: bool = False) -> bool:
        """Crea o actualiza el índice vectorial.

        Antes, en cada arranque del proceso se reindexaba y
        `Chroma.from_documents` AÑADÍA a la colección persistida, así que los
        fragmentos se duplicaban en cada reinicio. Ahora el fingerprint se
        guarda en disco y se reutiliza el índice si el contenido no cambió.
        Devuelve True si indexó desde cero.
        """
        fp_actual = self._calcular_fingerprint()

        # Reutilizar el índice persistido si el contenido no cambió
        if not force and fp_actual and self._leer_fingerprint_persistido() == fp_actual:
            self.cargar_documentos()
            if self.documents:
                os.makedirs(self.persist_dir, exist_ok=True)
                self.vectorstore = Chroma(
                    collection_name=self._nombre_coleccion(fp_actual),
                    embedding_function=self.embeddings,
                    persist_directory=self.persist_dir,
                )
                if self._tiene_datos():
                    self._fingerprint = fp_actual
                    return False
                # Fingerprint coincide pero la colección está vacía
                # (por ejemplo, borrado parcial): reconstruir.

        self.cargar_documentos()
        if not self.documents:
            self.vectorstore = None
            self._fingerprint = fp_actual
            self._escribir_fingerprint(fp_actual)
            return True

        splitter = RecursiveCharacterTextSplitter(
            chunk_size=CHUNK_SIZE,
            chunk_overlap=CHUNK_OVERLAP,
            separators=["\n\n", "\n", ". ", " ", ""],
        )
        chunks = splitter.split_documents(self.documents)
        os.makedirs(self.persist_dir, exist_ok=True)

        nombre = self._nombre_coleccion(fp_actual)
        self._borrar_coleccion(nombre)

        self.vectorstore = Chroma.from_documents(
            documents=chunks,
            embedding=self.embeddings,
            collection_name=nombre,
            persist_directory=self.persist_dir,
        )
        self._fingerprint = fp_actual
        self._escribir_fingerprint(fp_actual)
        return True

    def _borrar_coleccion(self, nombre: str) -> None:
        """Elimina la colección si existe, para reindexar sin duplicar."""
        try:
            viejo = Chroma(
                collection_name=nombre,
                embedding_function=self.embeddings,
                persist_directory=self.persist_dir,
            )
            viejo.delete_collection()
        except Exception:
            pass

    def _tiene_datos(self) -> bool:
        """True si la colección actual contiene al menos un fragmento."""
        if self.vectorstore is None:
            return False
        try:
            return bool(self.vectorstore.get(limit=1).get("ids"))
        except Exception:
            return True  # no se puede verificar: asumir que está bien

    # ----- Retrieval -----
    def recuperar(self, consulta: str, k: int = 4) -> list[str]:
        """Recupera los k fragmentos más relevantes."""
        if self.vectorstore is None:
            return []
        docs = self.vectorstore.similarity_search(consulta, k=k)
        return [d.page_content for d in docs]

    def esta_listo(self) -> bool:
        """True si hay documentos indexados y listos para retrieval."""
        return self.vectorstore is not None and len(self.documents) > 0


# ============================================================
# Gestor de asignaturas
# ============================================================
class GestorAsignaturas:
    """Escanea la carpeta asignaturas/ y lista los cursos disponibles."""

    @staticmethod
    def listar() -> list[str]:
        if not os.path.isdir(ASIGNATURAS_DIR):
            return []
        return sorted([
            d for d in os.listdir(ASIGNATURAS_DIR)
            if os.path.isdir(os.path.join(ASIGNATURAS_DIR, d))
            and not d.startswith(".")
        ])

    @staticmethod
    def nombre_legible(slug: str) -> str:
        return slug.replace("-", " ").title()

    @staticmethod
    def tiene_documentos(slug: str) -> bool:
        ruta = os.path.join(ASIGNATURAS_DIR, slug, "documentos")
        return os.path.isdir(ruta) and len(os.listdir(ruta)) > 0
